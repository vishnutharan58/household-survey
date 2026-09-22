from pptx import Presentation
import sys

def replace_text_in_shape(shape, replacements):
    if not shape.has_text_frame:
        return
    
    for paragraph in shape.text_frame.paragraphs:
        p_text = paragraph.text
        original_text = p_text
        for old, new in replacements.items():
            p_text = p_text.replace(old, new)
        
        if p_text != original_text:
            if len(paragraph.runs) > 0:
                font = paragraph.runs[0].font
                font_name = font.name if hasattr(font, 'name') else None
                font_size = font.size if hasattr(font, 'size') else None
                font_bold = font.bold if hasattr(font, 'bold') else None
                font_italic = font.italic if hasattr(font, 'italic') else None
                
                paragraph.text = p_text
                
                if len(paragraph.runs) > 0:
                    if font_name: paragraph.runs[0].font.name = font_name
                    if font_size: paragraph.runs[0].font.size = font_size
                    if font_bold is not None: paragraph.runs[0].font.bold = font_bold
                    if font_italic is not None: paragraph.runs[0].font.italic = font_italic
            else:
                paragraph.text = p_text

replacements = {
    "Green Voice Global": "Pro-Vision",
    "environmental NGO focused on afforestation, ecological restoration, and sustainability awareness among youth and local communities.": "public charitable trust empowering children, women and persons with disabilities through sustainable development.",
    "Velachery, Chennai": "Kanyakumari District",
    "Velachery": "Kanyakumari",
    "environment, health, and education since its formation in 2001.": "community development, empowering marginalized groups and conducting digital surveys.",
    "Mr. C.S. Veerraghavan": "the Board of Trustees",
    "living with visual impairment, driven by a personal passion for creating positive social change.": "dedicated to serving the marginalized and improving community living standards.",
    "His journey reflects that determination and vision matter more than circumstance": "Their journey reflects a deep commitment to sustainable, people-centric development",
    "He continues to personally engage": "The leadership continues to personally engage",
    "Environment, Education, Sustainable Development, and Health.": "Community Development, Digital Literacy, Health & Education.",
    "Environment: tree planting, sapling distribution, awareness campaigns, rainwater harvesting rallies": "Digital Surveying: Building a Household Survey app for accurate demographic data collection.",
    "This project engaged specifically with the Environment vertical.": "This project engaged specifically with the Digital Surveying vertical.",
    "its environment vertical offers direct, hands-on participation in tree plantation and awareness work": "its digital initiatives offered a chance to develop a real-world Household Survey application",
    "tree plantation techniques - pit preparation, spacing, soil treatment, and species selection suited to local conditions.": "app development techniques - React Native/Node.js, state management, form validation, and database architecture.",
    "poor soil quality, water scarcity, urban heat": "offline data sync, complex form handling, and user authentication",
    "assess sapling survival and basic aftercare requirements.": "test application stability and manage deployment requirements.",
    "Distributed environmental awareness booklets to local children in Velachery as part of Green Voice Global's community outreach program, encouraging early awareness of tree conservation.": "Week 1 & 2: Analyzed Pro-Vision's requirements for household data collection and designed the application architecture and database schema.",
    "Participated in a sapling plantation drive at Children's Park, Velachery - carefully planted and secured a young tree along the park boundary.": "Week 3 & 4: Developed the frontend UI for the Household Survey app, implementing complex forms for capturing demographic data.",
    "Came back a week later to check on the sapling's growth and make sure it was well-nourished.": "Week 5 & 6: Integrated the backend API with the database, ensuring secure storage of sensitive household information.",
    "Planted more saplings on a weekly basis": "Week 7 & 8: Conducted rigorous testing, fixed bugs, and demonstrated the completed Household Survey app to the organization.",
    "Environment Vertical (Primary Focus):": "Digital Surveying Vertical (Primary Focus):",
    "Tree planting drives in both rural and urban areas using locally-suited native species.": "Designing survey logic and offline forms to capture local demographic data efficiently.",
    "Sapling distribution to the public, along with aftercare guidance.": "Developing mobile and web interfaces for data entry workers to manage records.",
    "Awareness campaigns in schools and colleges on deforestation and climate change.": "Creating centralized databases to synchronize collected data seamlessly.",
    "Public rallies promoting rainwater harvesting and action on global warming.": "Implementing robust authentication systems to ensure data security and privacy.",
    "NGO": "Organization",
    "ngo": "organization",
    "environmental": "digital surveying",
    "afforestation": "software engineering",
    "tree conservation": "data collection",
    "sapling": "application module"
}

ppt_path = r"d:\HOUSEHOLDSURVEY\COMMUNITY CONNECT PPT FORMAT R.VishnuTharan.pptx"
out_path = r"d:\HOUSEHOLDSURVEY\COMMUNITY CONNECT PPT FORMAT R.VishnuTharan_Provision.pptx"

try:
    prs = Presentation(ppt_path)
    for slide in prs.slides:
        for shape in slide.shapes:
            replace_text_in_shape(shape, replacements)
            if shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        replace_text_in_shape(cell, replacements)

    prs.save(out_path)
    print(f"Successfully saved to {out_path}")
except Exception as e:
    print(f"Error: {e}")
