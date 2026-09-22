from pptx import Presentation
prs = Presentation('d:/HOUSEHOLDSURVEY/COMMUNITY CONNECT PPT FORMAT R.VishnuTharan_Provision.pptx')
slide = prs.slides[1] # Slide 2 is index 1

for shape in slide.shapes:
    if shape.has_text_frame:
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                if run.font.name == 'Noto Sans Symbols':
                    text = run.text
                    hex_val = hex(ord(text.strip())) if text.strip() else 'empty'
                    print(f"Symbol found: '{text}' (Unicode: {hex_val})")
